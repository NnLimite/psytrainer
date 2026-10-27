from __future__ import annotations

import io
import json
import zipfile

from app.payments import sign
from tests.conftest import login, make_user
from tests.test_assessment import RUBRIC, finished_session, run_assessment
from tests.test_orgs import setup
from tests.test_sessions import consent, ready, start


async def grouped(client, api, container):
    org, teacher, student, h, cid, vid = await ready(client, api, container)
    g = (await client.post(f"{api}/groups", json={"org_id": str(org.id), "name": "Г"}, headers=h["teacher"])).json()
    await client.put(f"{api}/groups/{g['id']}/members/{student.id}", headers=h["teacher"])
    rub = (await client.post(f"{api}/rubrics", json=RUBRIC, headers=h["teacher"])).json()
    return org, teacher, student, h, cid, vid, g, rub


async def test_analytics_respects_analytics_consent(client, api, container):
    org, teacher, student, h, cid, vid, g, rub = await grouped(client, api, container)
    s = await finished_session(client, api, h, vid, group_id=g["id"])
    await run_assessment(client, api, container, h["student"], s["id"], rub["id"])

    summary = (await client.get(f"{api}/analytics/groups/{g['id']}/summary", headers=h["teacher"])).json()
    assert summary["sessions_count"] == 0 and summary["avg_total_score"] is None  # нет согласия analytics
    await consent(client, api, h["student"], "analytics")
    summary = (await client.get(f"{api}/analytics/groups/{g['id']}/summary", headers=h["teacher"])).json()
    assert summary["sessions_count"] == 1 and summary["active_students"] == 1 and summary["avg_total_score"] == 75.0
    assert {x["code"] for x in summary["skills"]} == {"empathy", "questions"}
    url = f"{api}/analytics/groups/{g['id']}/summary"
    assert (await client.get(url, params={"from": "2999-01-01"}, headers=h["teacher"])).json()["sessions_count"] == 0
    assert (await client.get(url, params={"from": "2000-01-01", "to": "2999-01-01"}, headers=h["teacher"])).json()["sessions_count"] == 1

    usage = (await client.get(f"{api}/analytics/cases/{cid}/usage", headers=h["teacher"])).json()
    assert usage["sessions_count"] == 1 and usage["avg_final_trust"] > 0
    prog = (await client.get(f"{api}/analytics/students/{student.id}/progress", headers=h["teacher"])).json()
    assert prog["sessions_completed"] == 1 and prog["avg_total_score"] == 75.0
    mine = (await client.get(f"{api}/me/progress", headers=h["student"])).json()
    assert mine["sessions_completed"] == 1 and mine["skills"][0]["delta_30d"] is None
    assert (await client.get(f"{api}/me/progress", headers=h["teacher"])).status_code == 403
    assert (await client.get(f"{api}/analytics/groups/{g['id']}/summary", headers=h["student"])).status_code == 403
    stranger = await login(client, await make_user(container, "teacher", org))
    assert (await client.get(f"{api}/analytics/groups/{g['id']}/summary", headers=stranger)).status_code == 404
    assert (await client.get(f"{api}/analytics/students/{student.id}/progress", headers=stranger)).status_code == 404


async def test_preview_sessions_excluded_from_analytics(client, api, container):
    org, teacher, student, h, cid, vid, g, rub = await grouped(client, api, container)
    await consent(client, api, h["teacher"])
    await consent(client, api, h["student"], "analytics")
    await consent(client, api, h["teacher"], "analytics")
    p = (await start(client, api, h["teacher"], vid, mode="preview")).json()
    await client.post(f"{api}/sessions/{p['id']}/end", headers=h["teacher"])
    usage = (await client.get(f"{api}/analytics/cases/{cid}/usage", headers=h["teacher"])).json()
    assert usage["sessions_count"] == 0


async def test_reports_csv_xlsx_and_pdf_unsupported(client, api, container):
    org, teacher, student, h, cid, vid, g, rub = await grouped(client, api, container)
    await consent(client, api, h["student"], "analytics")
    s = await finished_session(client, api, h, vid, group_id=g["id"])
    await run_assessment(client, api, container, h["student"], s["id"], rub["id"])

    r = await client.post(f"{api}/reports", json={"type": "group_summary", "params": {"group_id": g["id"]}, "format": "csv"}, headers=h["teacher"])
    assert r.status_code == 202
    await container.jobs.drain()
    job = (await client.get(f"{api}/jobs/{r.json()['id']}", headers=h["teacher"])).json()
    assert job["status"] == "succeeded" and job["result_url"].endswith("/download")
    dl = await client.get(f"{api}/jobs/{job['id']}/download", headers=h["teacher"])
    assert dl.status_code == 200 and dl.headers["content-type"].startswith("text/csv")
    text = dl.content.decode("utf-8-sig")
    assert "Активных студентов,1" in text and "empathy" in text

    r = await client.post(f"{api}/reports", json={"type": "case_usage", "params": {"case_id": cid}, "format": "xlsx"}, headers=h["teacher"])
    await container.jobs.drain()
    dl = await client.get(f"{api}/jobs/{r.json()['id']}/download", headers=h["teacher"])
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(dl.content)).active
    assert ws["A1"].value == "Кейс" and ws["A2"].value == "Сессий" and ws["B2"].value == 1

    r = await client.post(f"{api}/reports", json={"type": "group_summary", "params": {"group_id": g["id"]}, "format": "pdf"}, headers=h["teacher"])
    assert r.status_code == 422 and r.json()["code"] == "format_not_supported"
    r = await client.post(f"{api}/reports", json={"type": "group_summary", "params": {}, "format": "csv"}, headers=h["teacher"])
    assert r.status_code == 422
    stranger = await login(client, await make_user(container, "teacher", org))
    r = await client.post(f"{api}/reports", json={"type": "group_summary", "params": {"group_id": g["id"]}, "format": "csv"}, headers=stranger)
    assert r.status_code == 404


async def test_data_export_and_erasure(client, api, container):
    org, teacher, student, h, cid, vid, g, rub = await grouped(client, api, container)
    s = await finished_session(client, api, h, vid, texts=("Здравствуйте", "Расскажите подробнее"), group_id=g["id"])
    ex = await client.post(f"{api}/me/data-export", headers=h["student"])
    assert ex.status_code == 202
    await container.jobs.drain()
    dl = await client.get(f"{api}/jobs/{ex.json()['id']}/download", headers=h["student"])
    data = json.loads(zipfile.ZipFile(io.BytesIO(dl.content)).read("export.json"))
    assert data["user"]["email"] == student.email and len(data["sessions"][0]["turns"]) == 4
    assert [c["kind"] for c in data["consents"]] == ["transcript_processing"]

    er = await client.delete(f"{api}/me", headers=h["student"])
    assert er.status_code == 202
    await container.jobs.drain()
    from app.models import Job
    async with container.session_factory() as db:
        assert (await db.get(Job, __import__("uuid").UUID(er.json()["id"]))).status == "succeeded"
    # после удаления войти нельзя, транскрипт стёрт, сессия осталась обезличенной
    r = await client.post(f"{api}/auth/login", json={"email": student.email, "password": "correct-horse-battery"})
    assert r.status_code == 401
    turns = (await client.get(f"{api}/sessions/{s['id']}/turns", headers=h["teacher"])).json()["items"]
    assert turns == []
    assert (await client.get(f"{api}/groups/{g['id']}/members", headers=h["teacher"])).json()["items"] == []


async def test_billing_checkout_and_webhook(client, api, container):
    org, admin, teacher, student, h = await setup(container, client)
    assert (await client.post(f"{api}/billing/checkout", json={"plan": "nope", "success_url": "https://x/s", "cancel_url": "https://x/c"}, headers=h["student"])).status_code == 422
    assert (await client.post(f"{api}/billing/checkout", json={"plan": "org_yearly", "success_url": "https://x/s", "cancel_url": "https://x/c"}, headers=h["teacher"])).status_code == 403
    r = await client.post(f"{api}/billing/checkout", json={"plan": "student_monthly", "success_url": "https://x/s", "cancel_url": "https://x/c"}, headers=h["student"])
    url = r.json()["checkout_url"]
    ref = url.split("/pay/stub/")[1].split("?")[0]

    assert (await client.get(f"{api}/me/subscription", headers=h["student"])).json()["plan"] == "free"
    body = json.dumps({"id": "evt-1", "type": "payment.succeeded", "data": {"ref": ref}}).encode()
    secret = container.settings.payment_webhook_secret
    assert (await client.post(f"{api}/webhooks/payments/stub", content=body, headers={"X-Signature": "bad"})).status_code == 401
    ok = {"X-Signature": sign(secret, body)}
    assert (await client.post(f"{api}/webhooks/payments/stub", content=body, headers=ok)).status_code == 204
    sub = (await client.get(f"{api}/me/subscription", headers=h["student"])).json()
    assert sub["plan"] == "student_monthly" and sub["status"] == "active" and sub["valid_until"]
    assert (await client.post(f"{api}/webhooks/payments/stub", content=body, headers=ok)).status_code == 204  # повтор не продлевает
    assert (await client.get(f"{api}/me/subscription", headers=h["student"])).json()["valid_until"] == sub["valid_until"]
    assert (await client.post(f"{api}/webhooks/payments/other", content=body, headers=ok)).status_code == 404

    # лицензия организации
    r = await client.post(f"{api}/billing/checkout", json={"plan": "org_yearly", "success_url": "https://x/s", "cancel_url": "https://x/c"}, headers=h["admin"])
    ref2 = r.json()["checkout_url"].split("/pay/stub/")[1].split("?")[0]
    body2 = json.dumps({"id": "evt-2", "type": "payment.succeeded", "data": {"ref": ref2}}).encode()
    await client.post(f"{api}/webhooks/payments/stub", content=body2, headers={"X-Signature": sign(secret, body2)})
    assert (await client.get(f"{api}/organizations/{org.id}/subscription", headers=h["admin"])).json()["plan"] == "org_yearly"
    assert (await client.get(f"{api}/me/subscription", headers=h["teacher"])).json()["plan"] == "org_yearly"  # лицензия действует на участников


async def test_admin_templates_and_llm_usage(client, api, container):
    org, teacher, student, h, cid, vid = await ready(client, api, container)
    admin = await login(client, await make_user(container, "platform_admin"))
    assert (await client.get(f"{api}/admin/prompt-templates", headers=h["teacher"])).status_code == 403
    bad = await client.post(f"{api}/admin/prompt-templates", json={"kind": "classifier", "body": "Без плейсхолдера, просто текст"}, headers=admin)
    assert bad.status_code == 422
    t1 = (await client.post(f"{api}/admin/prompt-templates", json={"kind": "classifier", "body": "Разметь реплику. Ключи: {keys}"}, headers=admin)).json()
    t2 = (await client.post(f"{api}/admin/prompt-templates", json={"kind": "classifier", "body": "Разметь точнее. Ключи: {keys}"}, headers=admin)).json()
    assert (t1["version"], t2["version"]) == (1, 2)
    await client.post(f"{api}/admin/prompt-templates/{t1['id']}/activate", headers=admin)
    await client.post(f"{api}/admin/prompt-templates/{t2['id']}/activate", headers=admin)
    items = (await client.get(f"{api}/admin/prompt-templates", params={"kind": "classifier"}, headers=admin)).json()["items"]
    assert [(i["version"], i["is_active"]) for i in items] == [(2, True), (1, False)]

    s = (await start(client, api, h["student"], vid)).json()
    await client.post(f"{api}/sessions/{s['id']}/messages", json={"text": "Привет"}, headers=h["student"])
    usage = (await client.get(f"{api}/admin/llm-usage", headers=admin)).json()
    by = {(i["model"], i["purpose"]): i for i in usage["items"]}
    assert by[("fake-qwen", "patient")]["calls"] == 1 and by[("fake-qwen", "classifier")]["tokens_in"] == 10
    assert by[("fake-qwen", "patient")]["p95_latency_ms"] == 7


async def test_system_and_error_shapes(client, api, container, llm):
    assert (await client.get(f"{api}/health")).json() == {"status": "ok", "checks": {}}
    r = await client.get(f"{api}/ready")
    assert r.status_code == 200 and r.json()["checks"] == {"database": "ok", "redis": "ok", "llm": "ok"}
    r = await client.get(f"{api}/no-such-route")
    assert r.status_code == 404 and r.json()["code"] == "not_found" and r.headers["x-request-id"]
    org, admin, teacher, student, h = await setup(container, client)
    r = await client.get(f"{api}/cases/not-a-uuid", headers=h["teacher"])
    assert r.status_code == 422 and r.json()["code"] == "validation_failed"
    r = await client.get(f"{api}/disorders", params={"cursor": "!!"}, headers=h["teacher"])
    assert r.status_code == 400
    assert (await client.get(f"{api}/disorders", params={"limit": 1000}, headers=h["teacher"])).status_code == 422


async def test_jobs_recover_after_restart(container):
    from app.models import Job
    uid = (await make_user(container)).id
    async with container.session_factory() as db:
        j = Job(user_id=uid, type="report", status="running")
        db.add(j)
        await db.commit()
        jid = j.id
    await container.jobs.recover()
    async with container.session_factory() as db:
        j = await db.get(Job, jid)
        assert j.status == "failed" and "restart" in j.error


async def test_purge_removes_old_job_files(container, tmp_path):
    from datetime import datetime, timedelta, timezone
    from app.models import Job
    uid = (await make_user(container)).id
    f = tmp_path / "old.csv"
    f.write_text("x")
    async with container.session_factory() as db:
        db.add_all([Job(user_id=uid, type="report", status="succeeded", result_file=str(f),
                        finished_at=datetime.now(timezone.utc) - timedelta(days=8)),
                    Job(user_id=uid, type="report", status="succeeded", finished_at=datetime.now(timezone.utc))])
        await db.commit()
    assert await container.jobs.purge(7) == 1
    assert not f.exists()
    assert await container.jobs.purge(7) == 0


async def test_cli_create_user_bootstraps_admin(container, client, api):
    from app.cli import create_user
    import pytest
    u = await create_user(container.session_factory, email="Boss@Uni.ru", password="long-enough-password", role="org_admin",
                          org_name="СПбГЭУ", sso_domain="unecon.ru")
    r = await client.post(f"{api}/auth/login", json={"email": "boss@uni.ru", "password": "long-enough-password"})
    tok = {"Authorization": f"Bearer {r.json()['access_token']}"}
    me = (await client.get(f"{api}/me", headers=tok)).json()
    assert me["role"] == "org_admin" and me["org_id"] == str(u.org_id)
    assert (await client.get(f"{api}/organizations/{u.org_id}", headers=tok)).json()["name"] == "СПбГЭУ"
    with pytest.raises(SystemExit):
        await create_user(container.session_factory, email="boss@uni.ru", password="long-enough-password", role="org_admin", org_name="СПбГЭУ")
    with pytest.raises(SystemExit):
        await create_user(container.session_factory, email="t@uni.ru", password="long-enough-password", role="teacher")


async def test_profile_patch_and_consents_listing(client, api, container):
    org, admin, teacher, student, h = await setup(container, client)
    r = await client.patch(f"{api}/me", json={"full_name": "Иван Петров", "locale": "en"}, headers=h["student"])
    assert r.status_code == 200 and r.json()["full_name"] == "Иван Петров" and r.json()["locale"] == "en"
    assert (await client.patch(f"{api}/me", json={"full_name": ""}, headers=h["student"])).status_code == 422
    assert (await client.get(f"{api}/me", headers=h["student"])).json()["full_name"] == "Иван Петров"

    lst = (await client.get(f"{api}/me/consents", headers=h["student"])).json()
    assert [(c["kind"], c["granted"]) for c in lst] == [("transcript_processing", False), ("analytics", False), ("research_use", False)]
    await consent(client, api, h["student"], "analytics")
    lst = (await client.get(f"{api}/me/consents", headers=h["student"])).json()
    assert {c["kind"]: c["granted"] for c in lst}["analytics"] is True
    await consent(client, api, h["student"], "analytics", granted=False)
    row = {c["kind"]: c for c in (await client.get(f"{api}/me/consents", headers=h["student"])).json()}["analytics"]
    assert row["granted"] is False and row["revoked_at"] is not None
    assert (await client.put(f"{api}/me/consents/unknown", json={"granted": True}, headers=h["student"])).status_code == 404


async def test_versions_listing_visibility(client, api, container):
    from tests.test_cases import create
    org, admin, teacher, student, h = await setup(container, client)
    cid = (await create(client, api, h["teacher"]))["id"]
    vs = (await client.get(f"{api}/cases/{cid}/versions", headers=h["teacher"])).json()
    assert [(v["version_no"], v["status"]) for v in vs] == [(1, "draft")]
    await client.post(f"{api}/cases/{cid}/versions/1/publish", headers=h["teacher"])
    await client.post(f"{api}/cases/{cid}/versions", json={}, headers=h["teacher"])  # v2: черновик
    mine = (await client.get(f"{api}/cases/{cid}/versions", headers=h["teacher"])).json()
    assert [(v["version_no"], v["status"]) for v in mine] == [(1, "published"), (2, "draft")]
    colleague = await login(client, await make_user(container, "teacher", org))
    theirs = (await client.get(f"{api}/cases/{cid}/versions", headers=colleague)).json()
    assert [(v["version_no"], v["status"]) for v in theirs] == [(1, "published")]  # чужой черновик не виден
    assert (await client.get(f"{api}/cases/{cid}/versions", headers=h["student"])).status_code == 403
